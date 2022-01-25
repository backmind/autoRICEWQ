# -*- coding: utf-8 -*-
"""
Created on Fri May 21 02:30:11 2021

@author: backm
"""
import pandas as pd
from io import StringIO
import os
import shutil
import re
import glob
import logging
from time import sleep


from openpyxl.chart import LineChart, Reference, Series
from openpyxl.chart.axis import DateAxis


def move_results(sim_name=""):
    if sim_name=="":
        sim_name="simulation"

    path = 'results\\'+sim_name
    files = glob.glob("bin\\*.*")
    files.remove('bin\\RICE192.EXE')
    try:
        files.remove('bin\\.gitignore')
    except:
        pass

    logging.shutdown()

    if os.path.exists(path):
        shutil.rmtree('results\\'+sim_name)

    os.makedirs('results\\'+sim_name)

    for file in files:
        os.rename(file, "results\\{0}\\{1}".format(sim_name, file.split("\\")[-1]))

def process_line(line_in):
    date_l = line_in[:10].split("/")[1].strip() + "-" + line_in[:10].split("/")[
        0].strip() + "-" + line_in[:10].split("/")[2].strip()
    data_l = re.sub(" +", ",", line_in[11:])
    line_out = date_l+data_l+"\n"
    return line_out


def process_pesticide_mass(sim_name="", delete_file=True, save_file=False):
    if sim_name == "":
        sim_name = "RICEWQ.ZP0"
    else:
        sim_name = sim_name + ".ZP0"

    with open("bin\\"+"RICEWQ.ZP0", "r") as text_file:
        file = text_file.read()

    file = file.split("Pesticide Mass (mg)\n")[-1]

    chems = []
    for chem in file.splitlines()[0].split()[1::2]:
        if '\x00' in chem:
            pass
        else:
            chems.append(chem)

    header1 = []
    for i, chem in enumerate(chems):
        header1.append("{0} - {1}".format(i+1, chem))
    header1.reverse()

    header2 = [
        "Pesticide in Water (mg)", "Pesticide in Sediment (mg)", "Pesticide in Foliage (mg)"]
    header2.reverse()

    header = ""
    for chem in header1:
        for word in header2:
            header = "{0}: {1}, {2}".format(chem, word, header)
    header = header[:-2] + "\n"

    out_file = ""
    for in_line in file.splitlines()[2:]:
        date_l = in_line[:10].split("/")[1].strip() + "-" + in_line[:10].split("/")[
            0].strip() + "-" + in_line[:10].split("/")[2].strip()
        data_l = in_line[11:].replace("   ", " ").replace(" ", ",")
        out_line = date_l+","+data_l+"\n"
        out_file = out_file+out_line

    out_file = header+out_file
    out_file = out_file[:-1]


    df = pd.read_csv(StringIO(out_file), header=[0])
    df.index = pd.to_datetime(df.index,dayfirst=True,format="%d-%m-%y",infer_datetime_format=True).date
    if save_file:
        df.to_excel("bin\\"+sim_name+".xlsx")
    if delete_file:
        os.remove("bin\\"+"RICEWQ.ZP0")
    return df


def process_paddy_outflow(sim_name="", delete_file=True, save_file=False):
    if sim_name == "":
        sim_name = "RICEWQ.ZZT"
    else:
        sim_name = sim_name + ".ZZT"

    with open("bin\\"+"RICEWQ.ZZT", "r") as text_file:
        file = text_file.read()

    file = file.split("(mg)\n")[-1]

    out_file = ""
    for line_in in file.splitlines()[1:]:
        line_out = process_line(line_in)
        out_file = out_file+line_out

    header = "JULIAN DAY,QOUT (m3),"
    for i in range(len(line_out.split(","))-3):
        header = header + "POUT%s (mg)," % str(i+1)
    header = header[:-1] + "\n"

    out_file = header+out_file
    out_file = out_file[:-1]

    df = pd.read_csv(StringIO(out_file), header=[0])
    df.index = pd.to_datetime(df.index,dayfirst=True,format="%d-%m-%y",infer_datetime_format=True).date
    if save_file:
        df.to_excel("bin\\"+sim_name+".xlsx")
    if delete_file:
        os.remove("bin\\"+"RICEWQ.ZZT")
    return df


def process_water_balance(sim_name="", delete_file=True, save_file=False):
    if sim_name == "":
        sim_name = "RICEWQ.ZZH"
    else:
        sim_name = sim_name + ".ZZH"

    with open("bin\\"+"RICEWQ.ZZH", "r") as text_file:
        file = text_file.read()

    file = file.split("(m^3)\n")[-1]

    out_file = ""
    for line_in in file.splitlines()[1:]:
        line_out = process_line(line_in)
        out_file = out_file+line_out

    header = "PRECIP (cm),EVAP (cm),SEEP (cm),SEEPS(cm),IRRIG (cm),IRRIG (cm),THETA (cm),DEPTH (cm),QOUT (m^3)\n"

    out_file = header+out_file
    out_file = out_file[:-1]

    df = pd.read_csv(StringIO(out_file), header=[0])
    df.index = pd.to_datetime(df.index,dayfirst=True,format="%d-%m-%y",infer_datetime_format=True).date
    if save_file:
        df.to_excel("bin\\"+sim_name+".xlsx")
    if delete_file:
        os.remove("bin\\"+"RICEWQ.ZZH")
    return df


def process_pesticide_balance(file_name, sim_name="", delete_file=True, save_file=False):
    if sim_name == "":
        sim_name = "RICEWQ.ZP"+file_name[-1]
    else:
        sim_name = sim_name + ".ZP"+file_name[-1]

    with open("bin\\"+file_name, "r") as text_file:
        file = text_file.read()
    try:
        file = file.split("(mg/kg)\n")[1]

        out_file = ""
        for line_in in file.splitlines()[1:]:
            line_out = process_line(line_in)
            out_file = out_file+line_out

        header = ['PWAP (mg)',
                  'PSAP (mg)',
                  'PFAP (mg)',
                  'PSSR (mg)',
                  'WO (mg)',
                  'DECAYW (mg)',
                  'DECAYS (mg)',
                  'DECAYF (mg)',
                  'VOLAT (mg)',
                  'SETL (mg)',
                  'BIND (mg)',
                  'SEEP (mg)',
                  'SEEPS (mg)',
                  'RESUS (mg)',
                  'DIFUS (mg)',
                  'PF (mg)',
                  'PW1 (mg)',
                  'PS1 (mg)',
                  'CPW (mg/l)',
                  'CPS (mg/kg)\n']

        if len(line_out.split(","))-1 == len(header):
            header = ",".join(header)
        else:
            header = ",".join(header[4:])

        out_file = header+out_file
        out_file = out_file[:-1]

        df = pd.read_csv(StringIO(out_file), header=[0])
        df.index = pd.to_datetime(df.index,dayfirst=True,format="%d-%m-%y",infer_datetime_format=True).date
        if save_file:
            df.to_excel("bin\\"+sim_name+".xlsx")
        if delete_file:
            os.remove("bin\\"+file_name)
    except:
        if delete_file:
            os.remove("bin\\"+file_name)
        df = pd.DataFrame()
    return df


def create_char(dates, data, title):
    chart = LineChart()

    chart.height = 10 # default is 7.5
    chart.width = 20 # default is 15

    chart.title = title
    chart.style = 13
    chart.x_axis.title = 'Date'
    # chart.y_axis.title = 'DEPTH (cm)'
    chart.y_axis.crossAx = 500
    chart.x_axis = DateAxis(crossAx=100)
    chart.x_axis.number_format ='dd/mm/yy'
    chart.x_axis.majorTimeUnit = "days"

    chart.add_data(data, titles_from_data=True)


    chart.set_categories(dates)
    return chart

def save_sim(sim_name="", delete_file=True, save_xlsxfiles=False, move=True, gen_summary=True):
    try:
        if sim_name == "":
            path = "simulation.xlsx"
        else:
            path = sim_name+".xlsx"
        writer = pd.ExcelWriter("bin\\"+path, engine='openpyxl', date_format = 'dd-mm-yy', datetime_format='dd-mm-yy')

        book = writer.book
        cs = book.create_sheet("Summary")

        df_water_balance = process_water_balance(sim_name, delete_file, save_xlsxfiles)
        df_water_balance.to_excel(writer, "water_balance (ZZH)")
        df_paddy_outflow = process_paddy_outflow(sim_name, delete_file, save_xlsxfiles)
        df_paddy_outflow.to_excel(writer, "paddy_outflow (ZZT)")
        df_pesticide_mass = process_pesticide_mass(
            sim_name, delete_file, save_xlsxfiles)
        df_pesticide_mass.to_excel(writer, "pesticide_mass (ZP0)")
        nchems=0
        for file_name in ["RICEWQ.ZP"+str(i+1) for i in range(5)]:
            df_pesticide_balance = process_pesticide_balance(
                file_name, sim_name, delete_file, save_xlsxfiles)
            if len(df_pesticide_balance) != 0:
                df_pesticide_balance.to_excel(
                    writer, "pesticide_mass ({0})".format(file_name.split(".")[-1]))
                nchems+=1

        if gen_summary:

            dates = Reference(writer.sheets["water_balance (ZZH)"], min_col=1, min_row=2, max_col=1, max_row=df_water_balance.shape[0]+1)


            data1 = Reference(writer.sheets["water_balance (ZZH)"], min_col=9, min_row=1, max_col=9, max_row=df_water_balance.shape[0]+1)
            data2 = Reference(writer.sheets["water_balance (ZZH)"], min_col=10, min_row=1, max_col=10, max_row=df_water_balance.shape[0]+1)
            data3 = Reference(writer.sheets["pesticide_mass (ZP1)"], min_col=20, min_row=1, max_col=20, max_row=df_water_balance.shape[0]+1)
            data4 = Reference(writer.sheets["pesticide_mass (ZP1)"], min_col=21, min_row=1, max_col=21, max_row=df_water_balance.shape[0]+1)

            cs.add_chart(create_char(dates, data1, "DEPTH"), "A1")
            cs.add_chart(create_char(dates, data2, "QOUT"), "M1")
            cs.add_chart(create_char(dates, data3, "CPW (mg/l)"), "Y1")
            cs.add_chart(create_char(dates, data4, "CPS (mg/kg)"), "AK1")

            namepos = 2
            for chemnumber in range(1,nchems+1):
                ypos = chemnumber*21
                title = writer.sheets["pesticide_mass (ZP0)"][1:1][chemnumber*3].value.split(":")[0]
                for ncol, col in enumerate(["A","M","Y"]):
                    data = Reference(writer.sheets["pesticide_mass (ZP0)"], min_col=namepos, min_row=1,  max_row=df_water_balance.shape[0]+1)
                    cs.add_chart(create_char(dates, data, title), col+str(ypos))
                    namepos += 1

            writer.save()
            writer.close()
    except:
        try:
            writer.close()
        except:
            pass

    if move:
        sleep(2)
        move_results(sim_name=sim_name)
